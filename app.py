from flask import Flask, render_template, request, jsonify
from datetime import datetime
import openpyxl
from openpyxl import load_workbook
import os

app = Flask(__name__)

EXCEL_FILE = os.path.join(os.path.dirname(__file__), "stars 3.xlsx")

# Noms forcés pour que ça marche tout de suite
NOMS_FORCES = ["Chloé", "Dorian", "Francky", "Jeremy", "Laurence", "Pouncho !"]

@app.route('/')
def index():
    return render_template('index.html', noms=NOMS_FORCES)

def trouver_colonne_jour(sheet, jour):
    for row in range(1, 25):
        for col in range(1, 400):
            try:
                val = sheet.cell(row=row, column=col).value
                if val and str(val).strip().isdigit() and int(float(val)) == jour:
                    return col
            except:
                continue
    return None

@app.route('/sauvegarder', methods=['POST'])
def sauvegarder():
    data = request.json
    nom = data['nom']
    date_str = data['date']
    selections = data['selections']

    try:
        date = datetime.strptime(date_str, "%Y-%m-%d")
        sheet_name = {1:"Janvier",2:"Février",3:"Mars",4:"Avril",5:"Mai",6:"Juin",
                      7:"Juillet",8:"Aout",9:"Septembre",10:"Octobre",11:"Novembre",12:"Decembre"}[date.month]
        jour = date.day
    except:
        return jsonify({"success": False, "message": "Date invalide"})

    wb = load_workbook(EXCEL_FILE)
    sheet = wb[sheet_name]

    # Recherche nom
    row_nom = None
    nom_clean = str(nom).strip().lower()
    for r in range(1, 100):
        for c in range(1, 6):
            cell = sheet.cell(row=r, column=c).value
            if cell and str(cell).strip().lower() == nom_clean:
                row_nom = r
                break
        if row_nom: break

    if not row_nom:
        return jsonify({"success": False, "message": f"Nom '{nom}' non trouvé"})

    col_m = trouver_colonne_jour(sheet, jour)
    if not col_m:
        return jsonify({"success": False, "message": f"Jour {jour} non trouvé"})

    for slot, choix in selections.items():
        valeur = "D" if choix == "DISPO" else "X"
        if slot == "Matin":
            sheet.cell(row=row_nom, column=col_m, value=valeur)
        elif slot == "Après-midi":
            sheet.cell(row=row_nom, column=col_m+1, value=valeur)
        elif slot == "Soir":
            sheet.cell(row=row_nom, column=col_m+2, value=valeur)

    wb.save(EXCEL_FILE)
    return jsonify({"success": True, "message": f"✅ Enregistré pour le {jour}/{date.month}"})

if __name__ == '__main__':
    print("🚀 Application démarrée avec noms forcés")
    app.run(host='0.0.0.0', port=5000, debug=True)
